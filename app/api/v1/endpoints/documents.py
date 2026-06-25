from __future__ import annotations

import io
import json
import uuid
from typing import AsyncGenerator, List

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user
from app.config import settings
from app.core.logging import get_logger
from app.database import AsyncSessionLocal, get_db
from app.models.document import Document, DocumentChunk
from app.models.user import User
from app.schemas.document import (
    DocumentCreate,
    DocumentListResponse,
    DocumentResponse,
    QueryRequest,
    QueryResponse,
)
from app.services.cache import CacheService, document_list_key, get_cache
from app.services.embedding import EmbeddingService, chunk_text, count_tokens
from app.services.rag import RAGService

router = APIRouter(prefix="/documents", tags=["Documents"])
logger = get_logger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _extract_text_from_pdf(data: bytes) -> str:
    """Extract all text from a PDF using pypdf."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages).strip()


def _extract_text_from_docx(data: bytes) -> str:
    """Extract all text from a DOCX file using python-docx."""
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs).strip()


async def _ingest_document(
    title: str,
    source: str | None,
    content: str,
    metadata: dict | None,
    cache: CacheService | None = None,
) -> Document:
    """
    Core ingestion logic: chunk → embed → persist.
    Invalidates the document-list cache after a successful write.
    """
    async with AsyncSessionLocal() as db:
        try:
            doc = Document(
                title=title,
                source=source,
                content=content,
                metadata_=json.dumps(metadata) if metadata else None,
            )
            db.add(doc)
            await db.flush()

            chunks = chunk_text(content)
            if not chunks:
                raise ValueError("Document content produced no chunks after tokenisation.")

            emb_service = EmbeddingService()
            embeddings = await emb_service.embed_batch(chunks, cache=cache)

            chunk_objs = [
                DocumentChunk(
                    document_id=doc.id,
                    chunk_index=i,
                    content=chunk,
                    embedding=emb,
                    token_count=count_tokens(chunk),
                )
                for i, (chunk, emb) in enumerate(zip(chunks, embeddings))
            ]
            db.add_all(chunk_objs)
            await db.commit()

            if cache:
                await cache.delete_pattern("rag:documents:list:*")

            logger.info(
                "Document ingested",
                document_id=str(doc.id),
                chunks=len(chunk_objs),
                title=title,
            )
            return doc
        except Exception:
            await db.rollback()
            raise


# ─────────────────────────────────────────────────────────────────────────────
# Ingest — JSON body
# ─────────────────────────────────────────────────────────────────────────────

@router.post(
    "/",
    response_model=DocumentResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Ingest a document (JSON)",
)
async def ingest_document(
    payload: DocumentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> DocumentResponse:
    doc = await _ingest_document(
        title=payload.title,
        source=payload.source,
        content=payload.content,
        metadata=payload.metadata,
        cache=cache,
    )
    return DocumentResponse(
        id=doc.id,
        title=doc.title,
        source=doc.source,
        content=doc.content,
        metadata=payload.metadata,
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


# ─────────────────────────────────────────────────────────────────────────────
# Ingest — File Upload (PDF / DOCX / TXT)
# ─────────────────────────────────────────────────────────────────────────────

@router.post(
    "/upload",
    response_model=DocumentResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Upload a file (PDF, DOCX, or TXT)",
)
async def upload_document(
    file: UploadFile = File(..., description="PDF, DOCX, or TXT file"),
    title: str | None = Query(None, description="Override document title (defaults to filename)"),
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> DocumentResponse:
    filename = file.filename or "uploaded_file"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in settings.allowed_extensions_list:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"File type '.{ext}' is not supported. Allowed: {settings.ALLOWED_EXTENSIONS}",
        )

    data = await file.read()
    if len(data) > settings.max_upload_size_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum size of {settings.MAX_UPLOAD_SIZE_MB} MB",
        )

    try:
        if ext == "pdf":
            content = _extract_text_from_pdf(data)
        elif ext == "docx":
            content = _extract_text_from_docx(data)
        else:
            content = data.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.error("File extraction failed", filename=filename, error=str(exc))
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Failed to extract text from file: {exc}",
        )

    if not content.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No text content could be extracted from the file.",
        )

    try:
        doc = await _ingest_document(
            title=title or filename,
            source=filename,
            content=content,
            metadata={"original_filename": filename, "file_size_bytes": len(data)},
            cache=cache,
        )
    except Exception as exc:
        err_str = str(exc)
        logger.error("Document ingestion failed", filename=filename, error=err_str, exc_info=True)
        if "401" in err_str or "authentication" in err_str.lower() or "api key" in err_str.lower():
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail=f"OpenAI authentication failed — check your OPENAI_API_KEY. Detail: {err_str}",
            )
        if "429" in err_str or "rate limit" in err_str.lower() or "quota" in err_str.lower():
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail=f"OpenAI rate limit / quota exceeded. Detail: {err_str}",
            )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Ingestion failed: {err_str}",
        )

    return DocumentResponse(
        id=doc.id,
        title=doc.title,
        source=doc.source,
        content=doc.content,
        metadata={"original_filename": filename, "file_size_bytes": len(data)},
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


# ─────────────────────────────────────────────────────────────────────────────
# Ingest — Async background
# ─────────────────────────────────────────────────────────────────────────────

@router.post(
    "/ingest-async",
    status_code=status.HTTP_202_ACCEPTED,
    summary="Ingest a document asynchronously (non-blocking)",
)
async def ingest_document_async(
    payload: DocumentCreate,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> dict:
    async def _background_job():
        try:
            await _ingest_document(
                title=payload.title,
                source=payload.source,
                content=payload.content,
                metadata=payload.metadata,
                cache=cache,
            )
        except Exception as exc:
            logger.error("Background ingestion failed", title=payload.title, error=str(exc))

    background_tasks.add_task(_background_job)
    return {
        "status": "accepted",
        "message": "Document is being processed in the background.",
        "title": payload.title,
    }


# ─────────────────────────────────────────────────────────────────────────────
# List
# ─────────────────────────────────────────────────────────────────────────────

@router.get(
    "/",
    response_model=DocumentListResponse,
    summary="List all documents (paginated)",
)
async def list_documents(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> DocumentListResponse:
    cache_key = document_list_key(page, page_size)

    cached = await cache.get(cache_key)
    if cached is not None:
        return DocumentListResponse(**cached)

    offset = (page - 1) * page_size
    total = (await db.execute(select(func.count(Document.id)))).scalar_one()
    docs = (
        await db.execute(
            select(Document)
            .order_by(Document.created_at.desc())
            .offset(offset)
            .limit(page_size)
        )
    ).scalars().all()

    result = DocumentListResponse(
        items=[
            DocumentResponse(
                id=d.id,
                title=d.title,
                source=d.source,
                content=d.content,
                metadata=json.loads(d.metadata_) if d.metadata_ else None,
                created_at=d.created_at,
                updated_at=d.updated_at,
            )
            for d in docs
        ],
        total=total,
        page=page,
        page_size=page_size,
    )

    await cache.set(cache_key, result.model_dump(mode="json"), ttl=settings.REDIS_TTL_DOCUMENT_LIST)
    return result


# ─────────────────────────────────────────────────────────────────────────────
# Get / Delete
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/{document_id}", response_model=DocumentResponse, summary="Get document by ID")
async def get_document(
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> DocumentResponse:
    doc = (
        await db.execute(select(Document).where(Document.id == document_id))
    ).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    return DocumentResponse(
        id=doc.id,
        title=doc.title,
        source=doc.source,
        content=doc.content,
        metadata=json.loads(doc.metadata_) if doc.metadata_ else None,
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


@router.delete(
    "/{document_id}",
    status_code=status.HTTP_200_OK,
    summary="Delete document and all its chunks",
)
async def delete_document(
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> dict:
    doc = (
        await db.execute(select(Document).where(Document.id == document_id))
    ).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    for chunk in (
        await db.execute(select(DocumentChunk).where(DocumentChunk.document_id == document_id))
    ).scalars().all():
        await db.delete(chunk)

    await db.delete(doc)
    await db.commit()
    await cache.delete_pattern("rag:documents:list:*")
    logger.info("Document deleted", document_id=str(document_id))
    return {"message": "Document deleted successfully", "id": str(document_id)}


# ─────────────────────────────────────────────────────────────────────────────
# RAG Query — Standard (blocking)
# ─────────────────────────────────────────────────────────────────────────────

@router.post(
    "/query",
    response_model=QueryResponse,
    summary="RAG query — blocking response",
)
async def rag_query(
    request: QueryRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    cache: CacheService = Depends(get_cache),
) -> QueryResponse:
    logger.info("RAG query", query=request.query, user=current_user.email)
    return await RAGService(db, cache=cache).query(request)


# ─────────────────────────────────────────────────────────────────────────────
# RAG Query — Streaming (Server-Sent Events)
# ─────────────────────────────────────────────────────────────────────────────

@router.post(
    "/query/stream",
    summary="RAG query — streaming response (SSE)",
    response_class=StreamingResponse,
)
async def rag_query_stream(
    request: QueryRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    logger.info("Streaming RAG query", query=request.query, user=current_user.email)

    async def _event_stream() -> AsyncGenerator[str, None]:
        from app.services.embedding import EmbeddingService, _is_mock_mode
        from openai import AsyncOpenAI

        try:
            emb_service = EmbeddingService()
            rag = RAGService(db)

            # 1. Embed query
            query_embedding = await emb_service.embed_text(request.query)

            # 2. Retrieve similar chunks
            rows = await rag.retrieve_similar_chunks(
                query_embedding=query_embedding,
                top_k=request.top_k,
                similarity_threshold=request.similarity_threshold,
            )

            # 3. Build context string
            if rows:
                context = "\n\n---\n\n".join(
                    f"[{i+1}] Source: {row.doc_title}\n{row.chunk_content}"
                    for i, row in enumerate(rows)
                )
            else:
                context = "No relevant context found."

            # 4. Emit sources event
            if request.include_sources:
                sources_payload = json.dumps([
                    {
                        "chunk_id": str(r.chunk_id),
                        "document_id": str(r.document_id),
                        "document_title": r.doc_title,
                        "similarity_score": round(r.similarity, 4),
                        "content": r.chunk_content,
                    }
                    for r in rows
                ])
                yield f"event: sources\ndata: {sources_payload}\n\n"

            # 5. Stream answer tokens — Ollama, OpenAI, or mock
            if _is_mock_mode():
                mock_answer = (
                    f"⚠️ Mock mode — no LLM configured.\n\n"
                    f"Found {len(rows)} relevant chunk(s). "
                    f"Set USE_OLLAMA=true or add OPENAI_API_KEY for real answers.\n\n"
                    f"Context preview:\n{context[:400]}..."
                )
                for word in mock_answer.split(" "):
                    yield f"data: {json.dumps({'delta': word + ' '})}\n\n"
                yield "data: [DONE]\n\n"
                return

            from app.services.embedding import _get_client
            client = _get_client()
            stream = await client.chat.completions.create(
                model=settings.chat_model,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You are a helpful AI assistant. Answer using ONLY the provided "
                            "context. If insufficient, say so. Cite sources as [1], [2], etc.\n\n"
                            f"Context:\n{context}"
                        ),
                    },
                    {"role": "user", "content": request.query},
                ],
                max_tokens=settings.OPENAI_MAX_TOKENS,
                temperature=settings.OPENAI_TEMPERATURE,
                stream=True,
            )

            async for chunk in stream:
                delta = chunk.choices[0].delta.content
                if delta:
                    yield f"data: {json.dumps({'delta': delta})}\n\n"

            yield "data: [DONE]\n\n"

        except Exception as exc:
            logger.error("SSE stream error", error=str(exc), exc_info=True)
            yield f"data: {json.dumps({'error': str(exc)})}\n\n"
            yield "data: [DONE]\n\n"

    return StreamingResponse(
        _event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )
